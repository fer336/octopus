"""
Tests de integración para exportación de reportes PDF, Excel y CSV.
"""

import io
from datetime import date, datetime
from decimal import Decimal
from uuid import uuid4

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.category import Category
from app.models.client import Client
from app.models.client_type import ClientType
from app.models.payment import Payment, PaymentMethod
from app.models.product import Product
from app.models.product_lot import ProductLot
from app.models.purchase_order import PurchaseOrder, PurchaseOrderItem, PurchaseOrderStatus
from app.models.stockpile import Stockpile
from app.models.supplier import Supplier
from app.models.user import User
from app.models.voucher import Voucher, VoucherStatus, VoucherType
from app.models.voucher_item import VoucherItem
from app.schemas.report_schemas import (
    CategoryReportFilters,
    ClientAccountsReportFilters,
    CurrentAccountWithdrawalsReportFilters,
    InventoryCountReportFilters,
    PurchaseOrderHistoryReportFilters,
    SalesReportFilters,
    StockReportFilters,
    StockpileWithdrawalsReportFilters,
    SupplierReportFilters,
    TopProductsReportFilters,
)
from app.services.reporting.category_report_service import CategoryReportService
from app.services.reporting.client_accounts_report_service import ClientAccountsReportService
from app.services.reporting.current_account_withdrawals_report_service import (
    CurrentAccountWithdrawalsReportService,
)
from app.services.reporting.inventory_count_report_service import InventoryCountReportService
from app.services.reporting.purchase_order_history_report_service import (
    PurchaseOrderHistoryReportService,
)
from app.services.reporting.sales_report_service import SalesReportService
from app.services.reporting.stock_report_service import StockReportService
from app.services.reporting.stockpile_withdrawals_report_service import (
    StockpileWithdrawalsReportService,
)
from app.services.reporting.supplier_report_service import SupplierReportService
from app.services.reporting.top_products_report_service import TopProductsReportService
from app.tests.conftest import make_auth_header


async def _create_category(db: AsyncSession, business_id, name: str = "Grifería") -> Category:
    """Crea una categoría mínima para reportes."""
    category = Category(business_id=business_id, name=name)
    db.add(category)
    await db.commit()
    await db.refresh(category)
    return category


async def _create_supplier(db: AsyncSession, business_id, name: str = "Proveedor Norte") -> Supplier:
    """Crea un proveedor mínimo para reportes."""
    supplier = Supplier(business_id=business_id, name=name, cuit="30-12345678-9")
    db.add(supplier)
    await db.commit()
    await db.refresh(supplier)
    return supplier


async def _create_contact_supplier(
    db: AsyncSession,
    business_id,
    *,
    name: str = "Proveedor Completo",
) -> Supplier:
    """Crea un proveedor con datos de contacto para la barra de filtros."""
    supplier = Supplier(
        business_id=business_id,
        name=name,
        cuit="30-87654321-0",
        phone="2222-3333",
        email="compras@example.com",
        contact_name="Ana Proveedor",
    )
    db.add(supplier)
    await db.commit()
    await db.refresh(supplier)
    return supplier


async def _create_product(
    db: AsyncSession,
    business_id,
    *,
    code: str = "PRD-001",
    description: str = "Producto test reporte",
    category_id=None,
    supplier_id=None,
    stock: int = 5,
    minimum_stock: int = 10,
    cost_price: Decimal = Decimal("100.00"),
    list_price: Decimal = Decimal("120.00"),
    discount_1: Decimal = Decimal("0.00"),
    discount_2: Decimal = Decimal("0.00"),
    discount_3: Decimal = Decimal("0.00"),
    sale_price: Decimal = Decimal("145.20"),
    is_active: bool = True,
) -> Product:
    """Crea un producto con lote para alimentar el stock calculado."""
    product = Product(
        business_id=business_id,
        category_id=category_id,
        supplier_id=supplier_id,
        code=code,
        supplier_code=f"SUP-{code}",
        description=description,
        minimum_stock=minimum_stock,
        cost_price=cost_price,
        list_price=list_price,
        discount_1=discount_1,
        discount_2=discount_2,
        discount_3=discount_3,
        discount_display="+".join(
            str(int(discount))
            for discount in [discount_1, discount_2, discount_3]
            if discount > 0
        ) or None,
        net_price=cost_price,
        sale_price=sale_price,
        iva_rate=Decimal("21.00"),
        is_active=is_active,
    )
    db.add(product)
    await db.commit()
    await db.refresh(product)

    lot = ProductLot(
        product_id=product.id,
        business_id=business_id,
        quantity=stock,
        initial_quantity=stock,
        received_date=date.today(),
    )
    db.add(lot)
    await db.commit()
    await db.refresh(product, attribute_names=["lots", "category", "supplier"])
    return product


async def _create_client_type(db: AsyncSession, business_id) -> ClientType:
    """Crea un tipo de cliente mínimo."""
    client_type = ClientType(business_id=business_id, name=f"Consumidor Final {uuid4()}")
    db.add(client_type)
    await db.commit()
    await db.refresh(client_type)
    return client_type


async def _create_client(
    db: AsyncSession,
    business_id,
    *,
    name: str = "Cliente Reporte",
    balance: Decimal = Decimal("0.00"),
) -> Client:
    """Crea un cliente mínimo con saldo configurable."""
    client_type = await _create_client_type(db, business_id)
    customer = Client(
        business_id=business_id,
        client_type_id=client_type.id,
        name=name,
        document_type="DNI",
        document_number="12345678",
        tax_condition="Consumidor Final",
        current_balance=balance,
        credit_limit=Decimal("50000.00"),
    )
    db.add(customer)
    await db.commit()
    await db.refresh(customer)
    return customer


async def _create_confirmed_voucher(
    db: AsyncSession,
    business_id,
    product: Product,
    client: Client,
    *,
    voucher_date: date = date(2026, 1, 15),
    voucher_type: VoucherType = VoucherType.INVOICE_B,
    number: str = "00000001",
    total: Decimal = Decimal("242.00"),
) -> Voucher:
    """Crea un comprobante confirmado con un ítem vendido."""
    voucher = Voucher(
        business_id=business_id,
        client_id=client.id,
        voucher_type=voucher_type,
        status=VoucherStatus.CONFIRMED,
        sale_point="0001",
        number=number,
        date=voucher_date,
        subtotal=Decimal("200.00"),
        iva_amount=Decimal("42.00"),
        total=total,
    )
    db.add(voucher)
    await db.commit()
    await db.refresh(voucher)

    item = VoucherItem(
        voucher_id=voucher.id,
        product_id=product.id,
        code=product.code,
        description=product.description,
        quantity=Decimal("2.00"),
        unit_price=Decimal("100.00"),
        iva_rate=Decimal("21.00"),
        iva_amount=Decimal("42.00"),
        subtotal=Decimal("200.00"),
        total=total,
        line_number=1,
    )
    db.add(item)
    await db.commit()
    return voucher


async def _create_purchase_order(
    db: AsyncSession,
    business_id,
    user: User,
    *,
    supplier_id=None,
    category_id=None,
    status: PurchaseOrderStatus = PurchaseOrderStatus.CONFIRMED,
    number: str = "00000001",
    created_at: datetime = datetime(2026, 1, 15, 10, 0),
) -> PurchaseOrder:
    """Crea una orden de pedido con un ítem y totales calculables."""
    product = await _create_product(
        db,
        business_id,
        code=f"PO-{number}",
        supplier_id=supplier_id,
        category_id=category_id,
        stock=4,
    )
    order = PurchaseOrder(
        business_id=business_id,
        supplier_id=supplier_id,
        category_id=category_id,
        created_by=user.id,
        status=status,
        sale_point="0001",
        number=number,
        subtotal=Decimal("300.00"),
        total_iva=Decimal("63.00"),
        total=Decimal("363.00"),
        confirmed_at=created_at if status == PurchaseOrderStatus.CONFIRMED else None,
        created_at=created_at,
    )
    db.add(order)
    await db.commit()
    await db.refresh(order)

    item = PurchaseOrderItem(
        purchase_order_id=order.id,
        product_id=product.id,
        system_stock=4,
        counted_stock=1,
        quantity_to_order=3,
        unit_cost=Decimal("100.00"),
        iva_rate=Decimal("21.00"),
        subtotal=Decimal("300.00"),
        iva_amount=Decimal("63.00"),
        total=Decimal("363.00"),
    )
    db.add(item)
    await db.commit()
    return order


async def _create_stockpile(db: AsyncSession, business_id, client: Client, name: str) -> Stockpile:
    """Crea un acopio mínimo para reportes de retiros."""
    stockpile = Stockpile(
        business_id=business_id,
        client_id=client.id,
        name=name,
        stockpile_number=f"AC-{uuid4().hex[:4].upper()}",
        initial_amount=Decimal("1000.00"),
        withdrawn_amount=Decimal("0.00"),
        remaining_amount=Decimal("1000.00"),
    )
    db.add(stockpile)
    await db.commit()
    await db.refresh(stockpile)
    return stockpile


def _xlsx_headers(content: bytes) -> list[str]:
    """Lee los encabezados del primer worksheet exportado."""
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    return [cell.value for cell in sheet[1]]


def _xlsx_data_row_count(content: bytes) -> int:
    """Cuenta filas de datos exportadas debajo de los encabezados."""
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    return max(sheet.max_row - 1, 0)


@pytest.mark.asyncio
async def test_stock_report_provider_builds_shared_dataset_with_totals(
    db: AsyncSession,
    business_a,
    user_a: User,
):
    """El proveedor de stock debe devolver filas y totales desde una única fuente."""
    category = await _create_category(db, business_a.id)
    supplier = await _create_supplier(db, business_a.id)
    await _create_product(
        db,
        business_a.id,
        category_id=category.id,
        supplier_id=supplier.id,
        stock=3,
        minimum_stock=5,
    )

    dataset = await StockReportService(db).build_dataset(
        business_a.id,
        StockReportFilters(low_stock_only=True),
        generated_by=user_a.email,
    )

    assert dataset.title == "Reporte de Stock"
    assert dataset.headers[:3] == ["Código", "Cód. Prov.", "Descripción"]
    assert dataset.rows[0]["Código"] == "PRD-001"
    assert dataset.rows[0]["Stock"] == 3
    assert dataset.totals["Ítems"] == 1
    assert dataset.totals["Stock bajo"] == 1


@pytest.mark.asyncio
async def test_stock_report_provider_returns_empty_dataset_for_no_matches(
    db: AsyncSession,
    business_a,
):
    """Los filtros sin resultados deben producir tabla vacía y totales cero."""
    await _create_product(db, business_a.id, description="No coincide")

    dataset = await StockReportService(db).build_dataset(
        business_a.id,
        StockReportFilters(search="inexistente"),
    )

    assert dataset.rows == []
    assert dataset.totals["Ítems"] == 0
    assert dataset.totals["Unidades"] == 0
    assert dataset.totals["Valor stock"] == 0.0


@pytest.mark.asyncio
async def test_sales_report_provider_preserves_date_and_receipt_filters(
    db: AsyncSession,
    business_a,
):
    """Ventas debe filtrar por período inclusivo y excluir remitos si se solicita."""
    product = await _create_product(db, business_a.id)
    customer = await _create_client(db, business_a.id)
    await _create_confirmed_voucher(db, business_a.id, product, customer, number="00000001")
    await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        customer,
        voucher_date=date(2026, 1, 16),
        voucher_type=VoucherType.RECEIPT,
        number="00000002",
    )

    dataset = await SalesReportService(db).build_dataset(
        business_a.id,
        SalesReportFilters(
            date_from=date(2026, 1, 1),
            date_to=date(2026, 1, 31),
            include_receipts=False,
        ),
    )

    assert [row["Tipo"] for row in dataset.rows] == ["invoice_b"]
    assert dataset.totals["Comprobantes"] == 1
    assert dataset.totals["Total"] == 242.0


@pytest.mark.asyncio
async def test_top_products_report_provider_respects_limit_and_order(
    db: AsyncSession,
    business_a,
):
    """Productos más vendidos debe respetar límite y ordenar por importe descendente."""
    customer = await _create_client(db, business_a.id)
    product_a = await _create_product(db, business_a.id, code="PRD-A", description="Producto A")
    product_b = await _create_product(db, business_a.id, code="PRD-B", description="Producto B")
    await _create_confirmed_voucher(db, business_a.id, product_a, customer, total=Decimal("242.00"))
    await _create_confirmed_voucher(
        db,
        business_a.id,
        product_b,
        customer,
        number="00000002",
        total=Decimal("484.00"),
    )

    dataset = await TopProductsReportService(db).build_dataset(
        business_a.id,
        TopProductsReportFilters(limit=1),
    )

    assert len(dataset.rows) == 1
    assert dataset.rows[0]["Código"] == "PRD-B"
    assert dataset.totals["Filas"] == 1
    assert dataset.totals["Cantidad"] == 2.0


@pytest.mark.asyncio
async def test_client_accounts_report_provider_filters_balances(
    db: AsyncSession,
    business_a,
):
    """Cuentas corrientes debe excluir clientes en cero cuando only_with_balance es true."""
    await _create_client(db, business_a.id, name="Cliente Deudor", balance=Decimal("1500.00"))
    await _create_client(db, business_a.id, name="Cliente Cero", balance=Decimal("0.00"))
    await _create_client(db, business_a.id, name="Cliente A Favor", balance=Decimal("-200.00"))

    dataset = await ClientAccountsReportService(db).build_dataset(
        business_a.id,
        ClientAccountsReportFilters(only_with_balance=True),
    )

    assert [row["Cliente"] for row in dataset.rows] == ["Cliente Deudor", "Cliente A Favor"]
    assert dataset.totals["Clientes"] == 2
    assert dataset.totals["Deuda"] == 1500.0
    assert dataset.totals["A favor"] == 200.0


@pytest.mark.asyncio
async def test_inventory_count_report_provider_reuses_count_sheet_rows(
    db: AsyncSession,
    business_a,
    user_a: User,
):
    """Conteo de inventario debe usar los productos activos de la planilla existente."""
    category = await _create_category(db, business_a.id)
    supplier = await _create_supplier(db, business_a.id)
    await _create_product(
        db,
        business_a.id,
        code="CNT-001",
        description="Producto para conteo",
        category_id=category.id,
        supplier_id=supplier.id,
        stock=8,
    )
    await _create_product(
        db,
        business_a.id,
        code="CNT-INACTIVE",
        category_id=category.id,
        supplier_id=supplier.id,
        is_active=False,
    )

    dataset = await InventoryCountReportService(db).build_dataset(
        business_a.id,
        InventoryCountReportFilters(supplier_id=supplier.id),
        generated_by=user_a.email,
    )

    assert dataset.title == "Planilla de Conteo de Inventario"
    assert dataset.headers == [
        "Código",
        "Descripción",
        "Categoría",
        "Proveedor",
        "Stock Sistema",
        "Conteo Físico",
        "A Pedir",
    ]
    assert dataset.rows == [
        {
            "Código": "CNT-001",
            "Descripción": "Producto para conteo",
            "Categoría": "Grifería",
            "Proveedor": "Proveedor Norte",
            "Stock Sistema": 8,
            "Conteo Físico": "",
            "A Pedir": "",
        }
    ]
    assert dataset.totals["Productos"] == 1
    assert dataset.orientation == "landscape"


@pytest.mark.asyncio
async def test_inventory_count_report_provider_returns_empty_dataset_for_valid_empty_filter(
    db: AsyncSession,
    business_a,
):
    """Un filtro válido sin productos debe generar tabla vacía con total cero."""
    supplier = await _create_supplier(db, business_a.id)

    dataset = await InventoryCountReportService(db).build_dataset(
        business_a.id,
        InventoryCountReportFilters(supplier_id=supplier.id),
    )

    assert dataset.rows == []
    assert dataset.totals["Productos"] == 0


@pytest.mark.asyncio
async def test_category_report_provider_groups_products_with_totals(
    db: AsyncSession,
    business_a,
):
    """Categorías debe agrupar productos y calcular totales por la fuente tabular."""
    category = await _create_category(db, business_a.id, name="Sanitarios")
    other_category = await _create_category(db, business_a.id, name="Grifería")
    supplier = await _create_supplier(db, business_a.id)
    await _create_product(
        db,
        business_a.id,
        code="CAT-001",
        description="Inodoro largo",
        category_id=category.id,
        supplier_id=supplier.id,
        stock=3,
        sale_price=Decimal("200.00"),
    )
    await _create_product(
        db,
        business_a.id,
        code="CAT-OUT",
        category_id=other_category.id,
        supplier_id=supplier.id,
    )

    dataset = await CategoryReportService(db).build_dataset(
        business_a.id,
        CategoryReportFilters(category_id=category.id),
    )

    assert dataset.title == "Reporte por Categoría"
    assert dataset.headers == [
        "Categoría",
        "Código",
        "Descripción",
        "Proveedor",
        "Stock",
        "Precio Venta",
        "Valor Stock",
    ]
    assert [row["Código"] for row in dataset.rows] == ["CAT-001"]
    assert dataset.rows[0]["Categoría"] == "Sanitarios"
    assert dataset.totals["Categorías"] == 1
    assert dataset.totals["Productos"] == 1
    assert dataset.totals["Unidades"] == 3
    assert dataset.totals["Valor stock"] == 600.0


@pytest.mark.asyncio
async def test_category_report_provider_returns_empty_dataset_for_no_matches(
    db: AsyncSession,
    business_a,
):
    """Categorías sin productos coincidentes debe devolver tabla vacía y totales cero."""
    category = await _create_category(db, business_a.id, name="Sin stock")

    dataset = await CategoryReportService(db).build_dataset(
        business_a.id,
        CategoryReportFilters(category_id=category.id),
    )

    assert dataset.rows == []
    assert dataset.totals["Categorías"] == 0
    assert dataset.totals["Productos"] == 0
    assert dataset.totals["Valor stock"] == 0.0


@pytest.mark.asyncio
async def test_supplier_report_provider_filters_identity_and_margin_totals(
    db: AsyncSession,
    business_a,
):
    """Proveedor debe filtrar productos, exponer identidad y calcular margen promedio."""
    category = await _create_category(db, business_a.id)
    supplier = await _create_contact_supplier(db, business_a.id)
    other_supplier = await _create_supplier(db, business_a.id, name="Proveedor Sur")
    await _create_product(
        db,
        business_a.id,
        code="SUP-001",
        description="Producto rentable",
        category_id=category.id,
        supplier_id=supplier.id,
        stock=5,
        cost_price=Decimal("100.00"),
        list_price=Decimal("140.00"),
        discount_1=Decimal("10.00"),
        sale_price=Decimal("150.00"),
    )
    await _create_product(
        db,
        business_a.id,
        code="SUP-OUT",
        supplier_id=other_supplier.id,
    )

    dataset = await SupplierReportService(db).build_dataset(
        business_a.id,
        SupplierReportFilters(supplier_id=supplier.id, only_with_stock=True),
    )

    assert dataset.title == "Reporte por Proveedor"
    assert dataset.filters["Proveedor"] == "Proveedor Completo"
    assert dataset.filters["CUIT"] == "30-87654321-0"
    assert "Ana Proveedor" in dataset.filters["Contacto"]
    assert [row["Código"] for row in dataset.rows] == ["SUP-001"]
    assert dataset.rows[0]["Bonificaciones"] == "10"
    assert dataset.rows[0]["Margen %"] == 50.0
    assert dataset.totals["Proveedores"] == 1
    assert dataset.totals["Productos"] == 1
    assert dataset.totals["Stock"] == 5
    assert dataset.totals["Valor stock"] == 500.0
    assert dataset.totals["Margen promedio"] == 50.0


@pytest.mark.asyncio
async def test_purchase_order_history_report_provider_filters_status_supplier_and_period(
    db: AsyncSession,
    business_a,
    user_a: User,
):
    """Historial de órdenes debe aplicar estado, proveedor y período inclusivo."""
    supplier = await _create_supplier(db, business_a.id)
    other_supplier = await _create_supplier(db, business_a.id, name="Proveedor Fuera")
    category = await _create_category(db, business_a.id)
    await _create_purchase_order(
        db,
        business_a.id,
        user_a,
        supplier_id=supplier.id,
        category_id=category.id,
        status=PurchaseOrderStatus.CONFIRMED,
        number="00000010",
        created_at=datetime(2026, 1, 20, 9, 0),
    )
    await _create_purchase_order(
        db,
        business_a.id,
        user_a,
        supplier_id=other_supplier.id,
        status=PurchaseOrderStatus.DRAFT,
        number="00000011",
        created_at=datetime(2026, 2, 20, 9, 0),
    )

    dataset = await PurchaseOrderHistoryReportService(db).build_dataset(
        business_a.id,
        PurchaseOrderHistoryReportFilters(
            date_from=date(2026, 1, 1),
            date_to=date(2026, 1, 31),
            supplier_id=supplier.id,
            status=PurchaseOrderStatus.CONFIRMED,
        ),
    )

    assert dataset.headers == [
        "Número",
        "Fecha",
        "Proveedor",
        "Categoría",
        "Ítems",
        "Estado",
        "Subtotal",
        "IVA",
        "Total",
    ]
    assert [row["Número"] for row in dataset.rows] == ["0001-00000010"]
    assert dataset.rows[0]["Estado"] == "confirmed"
    assert dataset.totals["Órdenes"] == 1
    assert dataset.totals["Ítems"] == 1
    assert dataset.totals["Subtotal"] == 300.0
    assert dataset.totals["IVA"] == 63.0
    assert dataset.totals["Total"] == 363.0


@pytest.mark.asyncio
async def test_stockpile_withdrawals_report_provider_includes_only_confirmed_stockpile_receipts(
    db: AsyncSession,
    business_a,
):
    """Retiros de acopio debe incluir remitos confirmados vinculados a acopios."""
    product = await _create_product(db, business_a.id, code="ACO-001")
    client = await _create_client(db, business_a.id, name="Cliente Acopio")
    stockpile = await _create_stockpile(db, business_a.id, client, "Obra Centro")
    included = await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        client,
        voucher_date=date(2026, 1, 10),
        voucher_type=VoucherType.RECEIPT,
        number="00000020",
    )
    included.stockpile_id = stockpile.id
    db.add(included)
    await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        client,
        voucher_date=date(2026, 1, 11),
        voucher_type=VoucherType.RECEIPT,
        number="00000021",
    )
    await db.commit()

    dataset = await StockpileWithdrawalsReportService(db).build_dataset(
        business_a.id,
        StockpileWithdrawalsReportFilters(
            date_from=date(2026, 1, 1),
            date_to=date(2026, 1, 31),
            stockpile_id=stockpile.id,
        ),
    )

    assert [row["Comprobante"] for row in dataset.rows] == ["0001-00000020"]
    assert dataset.rows[0]["Acopio"] == "Obra Centro"
    assert dataset.rows[0]["Cliente"] == "Cliente Acopio"
    assert dataset.totals["Filas"] == 1
    assert dataset.totals["Cantidad"] == 2.0
    assert dataset.totals["Valor"] == 242.0


@pytest.mark.asyncio
async def test_current_account_withdrawals_report_provider_excludes_non_withdrawal_movements(
    db: AsyncSession,
    business_a,
    user_a: User,
):
    """Retiros de cuenta corriente debe excluir pagos, devoluciones, cierres y facturas."""
    product = await _create_product(db, business_a.id, code="CC-001")
    billing_client = await _create_client(db, business_a.id, name="Cliente Facturación")
    withdrawal_client = await _create_client(db, business_a.id, name="Cliente Retira")
    valid = await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        billing_client,
        voucher_date=date(2026, 1, 10),
        voucher_type=VoucherType.RECEIPT,
        number="00000030",
    )
    valid.is_current_account = True
    valid.billing_client_id = billing_client.id
    valid.operating_client_id = withdrawal_client.id
    invoice = await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        billing_client,
        voucher_date=date(2026, 1, 11),
        voucher_type=VoucherType.INVOICE_B,
        number="00000031",
    )
    invoice.is_current_account = True
    closure = await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        billing_client,
        voucher_date=date(2026, 1, 12),
        voucher_type=VoucherType.RECEIPT,
        number="00000032",
    )
    closure.is_current_account = True
    closure.is_current_account_closure = True
    refund = await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        billing_client,
        voucher_date=date(2026, 1, 13),
        voucher_type=VoucherType.RECEIPT,
        number="00000033",
    )
    refund.is_current_account = True
    refund.is_return_receipt = True
    cash_receipt = await _create_confirmed_voucher(
        db,
        business_a.id,
        product,
        billing_client,
        voucher_date=date(2026, 1, 14),
        voucher_type=VoucherType.RECEIPT,
        number="00000034",
    )
    cash_receipt.is_current_account = False
    payment = Payment(
        business_id=business_a.id,
        client_id=billing_client.id,
        received_by=user_a.id,
        date=date(2026, 1, 15),
        amount=Decimal("100.00"),
        method=PaymentMethod.CASH,
    )
    db.add_all([valid, invoice, closure, refund, cash_receipt, payment])
    await db.commit()

    dataset = await CurrentAccountWithdrawalsReportService(db).build_dataset(
        business_a.id,
        CurrentAccountWithdrawalsReportFilters(client_id=billing_client.id),
    )

    assert [row["Comprobante"] for row in dataset.rows] == ["0001-00000030"]
    assert dataset.rows[0]["Cliente facturación"] == "Cliente Facturación"
    assert dataset.rows[0]["Cliente retiro"] == "Cliente Retira"
    assert dataset.totals["Filas"] == 1
    assert dataset.totals["Cantidad"] == 2.0
    assert dataset.totals["Subtotal"] == 200.0
    assert dataset.totals["IVA"] == 42.0
    assert dataset.totals["Total"] == 242.0


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("path", "expected_filename", "expected_header"),
    [
        ("/api/tenant/reports/stock", "reporte_stock_", "Código"),
        ("/api/tenant/reports/sales", "reporte_ventas_", "Fecha"),
        ("/api/tenant/reports/top-products", "reporte_productos_", "Código"),
        ("/api/tenant/reports/client-accounts", "reporte_cuentas_corrientes_", "Cliente"),
    ],
)
async def test_report_endpoints_export_pdf_xlsx_and_csv(
    client: AsyncClient,
    db: AsyncSession,
    user_a: User,
    business_a,
    membership_a,
    path: str,
    expected_filename: str,
    expected_header: str,
):
    """Cada reporte PR1 debe descargarse en PDF, Excel y CSV con el mismo contrato."""
    product = await _create_product(db, business_a.id)
    customer = await _create_client(db, business_a.id, balance=Decimal("100.00"))
    await _create_confirmed_voucher(db, business_a.id, product, customer)

    pdf_response = await client.get(path, headers=make_auth_header(user_a))
    assert pdf_response.status_code == 200, pdf_response.text
    assert pdf_response.headers["content-type"].startswith("application/pdf")
    assert expected_filename in pdf_response.headers.get("content-disposition", "")
    assert pdf_response.content.startswith(b"%PDF")

    xlsx_response = await client.get(f"{path}?format=xlsx", headers=make_auth_header(user_a))
    assert xlsx_response.status_code == 200, xlsx_response.text
    assert xlsx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert _xlsx_headers(xlsx_response.content)[0] == expected_header

    csv_response = await client.get(f"{path}?format=csv", headers=make_auth_header(user_a))
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert csv_response.text.splitlines()[0].startswith(expected_header)


@pytest.mark.asyncio
async def test_inventory_count_report_endpoint_exports_pdf_xlsx_and_csv(
    client: AsyncClient,
    db: AsyncSession,
    user_a: User,
    business_a,
    membership_a,
):
    """El hub de reportes debe descargar conteo de inventario en los tres formatos."""
    category = await _create_category(db, business_a.id)
    supplier = await _create_supplier(db, business_a.id)
    await _create_product(
        db,
        business_a.id,
        code="CNT-002",
        description="Producto exportable",
        category_id=category.id,
        supplier_id=supplier.id,
        stock=11,
    )
    path = f"/api/tenant/reports/inventory-count?supplier_id={supplier.id}"

    pdf_response = await client.get(path, headers=make_auth_header(user_a))
    assert pdf_response.status_code == 200, pdf_response.text
    assert pdf_response.headers["content-type"].startswith("application/pdf")
    assert "reporte_planilla_conteo_" in pdf_response.headers.get("content-disposition", "")
    assert pdf_response.content.startswith(b"%PDF")

    xlsx_response = await client.get(f"{path}&format=xlsx", headers=make_auth_header(user_a))
    assert xlsx_response.status_code == 200, xlsx_response.text
    assert xlsx_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert _xlsx_headers(xlsx_response.content) == [
        "Código",
        "Descripción",
        "Categoría",
        "Proveedor",
        "Stock Sistema",
        "Conteo Físico",
        "A Pedir",
    ]

    csv_response = await client.get(f"{path}&format=csv", headers=make_auth_header(user_a))
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.headers["content-type"].startswith("text/csv")
    lines = csv_response.text.splitlines()
    assert lines[0] == "Código,Descripción,Categoría,Proveedor,Stock Sistema,Conteo Físico,A Pedir"
    assert lines[1].startswith("CNT-002,Producto exportable,Grifería,Proveedor Norte,11,")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("path", "expected_filename", "expected_header"),
    [
        ("/api/tenant/reports/category", "reporte_categorias_", "Categoría"),
        ("/api/tenant/reports/supplier", "reporte_proveedores_", "Proveedor"),
        ("/api/tenant/reports/purchase-order-history", "reporte_historial_ordenes_", "Número"),
        ("/api/tenant/reports/stockpile-withdrawals", "reporte_retiros_acopio_", "Fecha"),
        ("/api/tenant/reports/current-account-withdrawals", "reporte_retiros_cuenta_corriente_", "Fecha"),
    ],
)
async def test_new_report_endpoints_export_empty_pdf_xlsx_and_csv(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
    path: str,
    expected_filename: str,
    expected_header: str,
):
    """Los cinco reportes nuevos deben descargar tablas vacías en los tres formatos."""
    pdf_response = await client.get(path, headers=make_auth_header(user_a))
    assert pdf_response.status_code == 200, pdf_response.text
    assert pdf_response.headers["content-type"].startswith("application/pdf")
    assert expected_filename in pdf_response.headers.get("content-disposition", "")
    assert pdf_response.content.startswith(b"%PDF")

    xlsx_response = await client.get(f"{path}?format=xlsx", headers=make_auth_header(user_a))
    assert xlsx_response.status_code == 200, xlsx_response.text
    assert _xlsx_headers(xlsx_response.content)[0] == expected_header

    csv_response = await client.get(f"{path}?format=csv", headers=make_auth_header(user_a))
    assert csv_response.status_code == 200, csv_response.text
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert csv_response.text.splitlines()[0].startswith(expected_header)


@pytest.mark.asyncio
async def test_new_report_unknown_filter_returns_400(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
):
    """Los reportes nuevos deben rechazar filtros no declarados."""
    response = await client.get(
        "/api/tenant/reports/current-account-withdrawals?not_allowed=1",
        headers=make_auth_header(user_a),
    )

    assert response.status_code == 400
    assert "Filtro no permitido" in response.text


@pytest.mark.asyncio
async def test_new_report_invalid_date_range_returns_400(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
):
    """Los reportes nuevos deben validar rangos de fechas antes de generar archivos."""
    response = await client.get(
        "/api/tenant/reports/purchase-order-history?date_from=2026-02-01&date_to=2026-01-01",
        headers=make_auth_header(user_a),
    )

    assert response.status_code == 400
    assert not response.content.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_inventory_count_report_requires_supplier_or_category_filter(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
):
    """La ruta de reportes debe rechazar planillas de conteo sin filtros."""
    response = await client.get(
        "/api/tenant/reports/inventory-count",
        headers=make_auth_header(user_a),
    )

    assert response.status_code == 400
    assert "proveedor o una categoría" in response.text


@pytest.mark.asyncio
async def test_compatibility_stock_pdf_alias_still_works(
    client: AsyncClient,
    db: AsyncSession,
    user_a: User,
    business_a,
    membership_a,
):
    """El alias histórico /stock/pdf debe seguir disponible para el frontend existente."""
    await _create_product(db, business_a.id)

    response = await client.get(
        "/api/tenant/reports/stock/pdf",
        headers=make_auth_header(user_a),
    )

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.content.startswith(b"%PDF")


@pytest.mark.asyncio
async def test_unknown_report_filter_returns_400(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
):
    """Los filtros desconocidos deben rechazarse antes de generar archivos."""
    response = await client.get(
        "/api/tenant/reports/sales?unknown_filter=1",
        headers=make_auth_header(user_a),
    )

    assert response.status_code == 400
    assert "Filtro no permitido" in response.text


@pytest.mark.asyncio
async def test_invalid_report_date_range_returns_400(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
):
    """Un rango de fechas inválido debe devolver HTTP 400 sin archivo."""
    response = await client.get(
        "/api/tenant/reports/sales?date_from=2026-02-01&date_to=2026-01-01",
        headers=make_auth_header(user_a),
    )

    assert response.status_code == 400
    assert not response.content.startswith(b"%PDF")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "path",
    [
        "/api/tenant/reports/sales?unknown_filter=1",
        "/api/tenant/reports/category?unexpected=1",
        "/api/tenant/reports/current-account-withdrawals?not_allowed=1",
    ],
)
async def test_cross_cutting_unknown_filters_return_400_before_file_generation(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
    path: str,
):
    """Los filtros desconocidos representativos deben fallar antes de renderizar archivos."""
    response = await client.get(path, headers=make_auth_header(user_a))

    assert response.status_code == 400
    assert "Filtro no permitido" in response.text
    assert not response.content.startswith(b"%PDF")


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "path",
    [
        "/api/tenant/reports/sales?date_from=2026-02-01&date_to=2026-01-01",
        "/api/tenant/reports/purchase-order-history?date_from=2026-02-01&date_to=2026-01-01",
        "/api/tenant/reports/stockpile-withdrawals?date_from=2026-02-01&date_to=2026-01-01",
    ],
)
async def test_cross_cutting_invalid_date_ranges_return_400_before_file_generation(
    client: AsyncClient,
    user_a: User,
    business_a,
    membership_a,
    path: str,
):
    """Los rangos inválidos representativos deben fallar sin contenido descargable."""
    response = await client.get(path, headers=make_auth_header(user_a))

    assert response.status_code == 400
    assert not response.content.startswith(b"%PDF")
    assert not response.headers.get("content-type", "").startswith("application/pdf")


@pytest.mark.asyncio
async def test_all_report_endpoints_return_empty_downloadable_tables_for_valid_empty_results(
    client: AsyncClient,
    db: AsyncSession,
    user_a: User,
    business_a,
    membership_a,
):
    """Todos los reportes deben exportar una tabla vacía válida cuando no hay datos."""
    empty_category = await _create_category(db, business_a.id, name="Categoría sin productos")
    empty_supplier = await _create_supplier(db, business_a.id, name="Proveedor sin productos")
    empty_stockpile_client = await _create_client(db, business_a.id, name="Cliente sin retiros")
    empty_stockpile = await _create_stockpile(
        db,
        business_a.id,
        empty_stockpile_client,
        "Acopio sin retiros",
    )

    report_cases = [
        ("/api/tenant/reports/stock?search=sin-resultados", "Código"),
        ("/api/tenant/reports/sales?date_from=2030-01-01&date_to=2030-01-31", "Fecha"),
        ("/api/tenant/reports/top-products?date_from=2030-01-01&date_to=2030-01-31", "Código"),
        ("/api/tenant/reports/client-accounts?only_with_balance=true", "Cliente"),
        (f"/api/tenant/reports/inventory-count?supplier_id={empty_supplier.id}", "Código"),
        (f"/api/tenant/reports/category?category_id={empty_category.id}", "Categoría"),
        (f"/api/tenant/reports/supplier?supplier_id={empty_supplier.id}", "Proveedor"),
        ("/api/tenant/reports/purchase-order-history?date_from=2030-01-01&date_to=2030-01-31", "Número"),
        (f"/api/tenant/reports/stockpile-withdrawals?stockpile_id={empty_stockpile.id}", "Fecha"),
        (
            f"/api/tenant/reports/current-account-withdrawals?client_id={empty_stockpile_client.id}",
            "Fecha",
        ),
    ]

    for path, expected_header in report_cases:
        separator = "&" if "?" in path else "?"

        pdf_response = await client.get(path, headers=make_auth_header(user_a))
        assert pdf_response.status_code == 200, f"{path}: {pdf_response.text}"
        assert pdf_response.headers["content-type"].startswith("application/pdf")
        assert pdf_response.content.startswith(b"%PDF")

        xlsx_response = await client.get(
            f"{path}{separator}format=xlsx",
            headers=make_auth_header(user_a),
        )
        assert xlsx_response.status_code == 200, f"{path}: {xlsx_response.text}"
        assert _xlsx_headers(xlsx_response.content)[0] == expected_header
        assert _xlsx_data_row_count(xlsx_response.content) == 0

        csv_response = await client.get(
            f"{path}{separator}format=csv",
            headers=make_auth_header(user_a),
        )
        assert csv_response.status_code == 200, f"{path}: {csv_response.text}"
        assert csv_response.headers["content-type"].startswith("text/csv")
        csv_lines = csv_response.text.splitlines()
        assert csv_lines[0].startswith(expected_header)
        assert len(csv_lines) == 1
